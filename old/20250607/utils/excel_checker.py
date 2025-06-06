import os
import glob
import logging
from openpyxl import load_workbook
from config.paths import REQUIRED_FILES
from utils.log_writer import save_log
from utils.template_loader import (
    load_check_template_paper1,
    load_check_template_paper2_1,
    load_check_template_paper2_2,
    load_check_template_paper3
)
from utils.checklist_writer import copy_checklist_template_for_club

def check_club_info(ws, club_master_dict):
    club_name = ws["A1"].value
    application_type = ws["A3"].value
    messages = []

    if not club_name:
        messages.append("❗ クラブ名が未記入です。")
    elif club_name not in club_master_dict:
        messages.append(f"❗ クラブ名『{club_name}』は名簿に存在しません。")
    else:
        expected = club_master_dict[club_name]
        if application_type != expected:
            messages.append(f"⚠️ 申請区分が不一致：名簿では『{expected}』、申請書では『{application_type}』です。")

    return messages

def check_cells_with_template(ws, template_rows, doc_name=""):
    prefix = f"[{doc_name}] " if doc_name else ""
    issues = []

    for item in template_rows:
        cell = item.get("セル位置")
        name = item.get("項目名")
        check_type = item.get("チェック種別")
        required = item.get("必須か")
        notes = item.get("備考", "")

        # 範囲指定セルの処理
        if ":" in cell:
            try:
                start_cell, end_cell = cell.split(":")
                for row in ws[start_cell:end_cell]:
                    for c in row:
                        val = c.value
                        if check_type == "○の有無" and required == "○":
                            valid_marks = {"○", "有", "1", "✔", "yes", "Yes"}
                            if not val or str(val).strip() not in valid_marks:
                                issues.append(f"{prefix}❗ {name}（{c.coordinate}）に○印がありません。")
            except Exception as e:
                issues.append(f"{prefix}❌ セル範囲 {cell} のチェック中にエラーが発生しました。")
            continue  # 個別セル処理をスキップ

        value = ws[cell].value if cell else None

        if check_type == "空欄チェック" and required == "○":
            if not value or str(value).strip() == "":
                issues.append(f"❗ {prefix}❗ {name}（{cell}）が未記入です。")

        elif check_type == "○の有無" and required == "○":
            if not value or "○" not in str(value):
                issues.append(f"❗ {prefix}❗ {name}（{cell}）に○印がありません。")

        elif check_type == "値チェック" and required == "○":
            if str(value).strip() not in ["新規", "更新"]:
                issues.append(f"❗ {prefix}❗ {name}（{cell}）が「新規／更新」になっていません。")

        elif check_type == "形式チェック":
            if "電話" in name and (not value or "-" not in str(value)):
                issues.append(f"{prefix}❗ {name}（{cell}）の形式が不正です。")
            elif "メール" in name and (not value or "@" not in str(value)):
                issues.append(f"{prefix}❗ {name}（{cell}）が正しいメールアドレス形式ではありません。")
            elif "人数" in name:
                try:
                    num = int(value)
                    if num < 0:
                        issues.append(f"{prefix}❗ {name}（{cell}）が0以上の数値ではありません。")
                except:
                    issues.append(f"{prefix}❗ {name}（{cell}）が数値として読み取れません。")

    return issues

def check_excel_contents_paper1(club_path, club_name, club_master_dict):
    issues = []

    pattern = os.path.join(club_path, "01_*確認用紙.xlsx")
    files = glob.glob(pattern)

    if not files:
        issues.append("❗ 登録基準確認用紙が見つかりません。")
        return issues

    file_path = files[0]
    try:
        wb = load_workbook(file_path)
        ws = wb.active

        # 名簿との照合チェック（A1, A3）
        issues += check_club_info(ws, club_master_dict)

    except Exception as e:
        logging.exception(f"{club_name}: Excel内容の読み取り中にエラー")
        issues.append("❌ Excelファイルの読み取りに失敗しました。")

    return issues

def check_excel_contents_paper2_1(club_path, club_name):
    issues = []

    pattern = os.path.join(club_path, "02-1_*基礎情報書類.xlsx")
    files = glob.glob(pattern)

    if not files:
        issues.append("❗ 書類②-1『基礎情報書類』が見つかりません。")
        return issues

    file_path = files[0]
    try:
        wb = load_workbook(file_path)
        ws = wb.active

    except Exception as e:
        logging.exception(f"{club_name}: 書類②-1のExcel読み込みエラー")
        issues.append("❌ 書類②-1の読み取りに失敗しました。")

    return issues

def check_excel_contents_paper2_2(club_path, club_name):
    issues = []

    pattern = os.path.join(club_path, "02-2_*.xlsx")
    files = glob.glob(pattern)

    if not files:
        issues.append("❗ 書類②-2『活動・マネジャー書類』が見つかりません。")
        return issues

    file_path = files[0]
    try:
        wb = load_workbook(file_path)
        ws = wb.active

    except Exception as e:
        logging.exception(f"{club_name}: 書類②-2のExcel読み込みエラー")
        issues.append("❌ 書類②-2の読み取りに失敗しました。")

    return issues

def check_excel_contents_paper3(club_path, club_name):
    issues = []

    pattern = os.path.join(club_path, "03_*.pdf")
    files = glob.glob(pattern)

    if not files:
        issues.append("❗ 書類③『規約・会則・定款等』が見つかりません。")
        return issues

    # 自動チェック処理は現時点で行わない（テンプレート確認も削除）
    return issues

def check_submissions(clubs, club_master):
    from config.paths import REQUIRED_FILES  # 局所的にも使えるよう再import（念のため）
    results = []

    for club_name, club_path in clubs:
        row = {"クラブ名": club_name}
        log_messages = []

        try:
            # 書類の存在チェック
            for code, filename in REQUIRED_FILES.items():
                pattern = os.path.join(club_path, f"{code}_*")
                files = glob.glob(pattern)
                if files:
                    row[filename] = "✔"
                else:
                    row[filename] = "×"
                    log_messages.append(f"❗ {filename} が見つかりません。")

            copy_checklist_template_for_club(club_name)

            # Excelの中身チェック
            if row[REQUIRED_FILES["01"]] == "✔":
                log_messages += check_excel_contents_paper1(club_path, club_name, club_master)
            
            if row.get(REQUIRED_FILES["02-1"]) == "✔":
                log_messages += check_excel_contents_paper2_1(club_path, club_name)
            
            if row.get(REQUIRED_FILES["02-2"]) == "✔":
                log_messages += check_excel_contents_paper2_2(club_path, club_name)
            
            if row.get(REQUIRED_FILES["03"]) == "✔":
                log_messages += check_excel_contents_paper3(club_path, club_name)

            # クラブ別ログ出力
            save_log(club_name, row, log_messages)
            logging.info(f"{club_name}: チェック完了")

        except Exception as e:
            logging.exception(f"{club_name}: 処理中にエラーが発生しました")

        results.append(row)

    return results
